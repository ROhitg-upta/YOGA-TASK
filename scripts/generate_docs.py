import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Set page margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base styling
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Calibri'
    normal_font.size = Pt(11)
    normal_font.color.rgb = RGBColor(0x2A, 0x2B, 0x28)

    # Document Header / Pre-title
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(4)
    run_meta = p_meta.add_run("STUDENT YOGIC CLUB (SYC) • ABES ENGINEERING COLLEGE • COHORT 2026–27")
    run_meta.font.name = 'Calibri'
    run_meta.font.size = Pt(9.5)
    run_meta.font.bold = True
    run_meta.font.color.rgb = RGBColor(0x2D, 0x4A, 0x3E) # Forest Green

    # Main Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    run_title = p_title.add_run("CONTENT LEAD — OFFICIAL INTERVIEW TASK SUBMISSION")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1C, 0x1D, 0x1A)

    # Subtitle / Dossier Info
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("Editorial Strategy, High-Conversion Social Copywriting & 20-30s Reel Production Script\n")
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x52, 0x50, 0x4A)
    
    run_deadline = p_sub.add_run("Domain: Content & Editorial Wing  |  Target Cohort: 2026–27  |  Status: Complete & Ready for Review")
    run_deadline.font.size = Pt(9.5)
    run_deadline.font.bold = True
    run_deadline.font.color.rgb = RGBColor(0x78, 0x76, 0x70)

    # Divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(14)
    p_div_run = p_div.add_run("―" * 58)
    p_div_run.font.color.rgb = RGBColor(0xDC, 0xD6, 0xC8)

    # Helper function for Section Headings
    def add_section_heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.bold = True
        if level == 1:
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(0x2D, 0x4A, 0x3E)
        elif level == 2:
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x1C, 0x1D, 0x1A)
        else:
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(0x3E, 0x3C, 0x36)
        return p

    # Helper function for Callout Box
    def add_callout(title, body_lines, tag="KEY TAKEAWAY", bg_color="F5F2EB"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(4)
        run_tag = p.add_run(f"[{tag}]  {title}\n")
        run_tag.font.bold = True
        run_tag.font.size = Pt(10)
        run_tag.font.color.rgb = RGBColor(0x2D, 0x4A, 0x3E)

        for line in body_lines:
            p_body = cell.add_paragraph()
            p_body.paragraph_format.space_after = Pt(3)
            r = p_body.add_run(line)
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(0x2A, 0x2B, 0x28)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # SECTION 1: EXECUTIVE EDITORIAL PHILOSOPHY
    add_section_heading("EXECUTIVE EDITORIAL PHILOSOPHY: MAKING YOGA UNIGNORABLE", level=1)
    
    p = doc.add_paragraph(
        "The Student Yogic Club (SYC) occupies a rare, high-leverage position at ABES Engineering College: "
        "it is the single campus society that harmonizes mental poise and inner discipline with elite engineering craftsmanship, "
        "visual design, and cinematic storytelling."
    )
    p.paragraph_format.space_after = Pt(6)

    p2 = doc.add_paragraph(
        "Most collegiate content fails because it either falls into slow, preachy spiritual clichés (\"Find your inner peace\"), "
        "or sounds like a dry notice board circular. My editorial mandate bridges the visceral reality of engineering students: "
        "8:30 AM attendance alarms, 2 AM DSA anxiety, and screen fatigue, positioning SYC as the ultimate cognitive and creative superpower."
    )
    p2.paragraph_format.space_after = Pt(12)

    # SECTION 2: INSTAGRAM CAPTIONS (6 Captions)
    add_section_heading("PART 1: INSTAGRAM CAPTIONS (3 POSTS × 2 HIGH-CONVERTING ANGLES)", level=1)

    # Post 1
    add_section_heading("Post 1: Sunrise Asana & Pranayama on the Campus Lawn", level=2)
    doc.add_paragraph(
        "Visual Context: High-aesthetic golden hour photo of students seated in Sukhasana on the morning grass, breathing in sync amidst morning mist."
    ).paragraph_format.space_after = Pt(6)

    add_callout(
        "OPTION 1A — The Poetic & Sensory Angle (The 'Campus Pause' Hook)",
        [
            "\"While the rest of campus is hitting snooze for the fourth time, something quiet is happening on the back lawn.\"",
            "",
            "8 hours of lectures. 3 assignment deadlines. Endless blue-light scrolling before bed.",
            "College moves fast, but your nervous system wasn’t built to live in a perpetual state of emergency.",
            "",
            "This morning, 45 of us traded notification alerts for the rhythm of breath. No laptops, no pending lab records, no expectations—just thirty minutes of grounding your mind before the world demands your attention.",
            "",
            "You don't need two hours of meditation to feel human again. You just need five intentional minutes of silence.",
            "",
            "🌿 CTA: If your morning routine currently starts with opening Instagram before getting out of bed, save this post as your reminder to step outside tomorrow morning. Tell us in the comments: What’s the first thing you usually do when you wake up?",
            "",
            "Hashtags: #StudentYogicClub #SYCABES #MorningFlow #CampusMindfulness #EngineeringSanctuary #BreatheBeforeYouBuild #CollegeWellbeing #PranayamaMornings"
        ],
        tag="CAPTION 1A",
        bg_color="EAF2EC"
    )

    add_callout(
        "OPTION 1B — The High-Energy, Student-Relatable Angle (The 'Caffeine Crash' Hook)",
        [
            "\"The most underrated cheat code for a 9-to-5 college schedule isn't a third cup of canteen coffee. It’s this.\"",
            "",
            "Let’s be real: college burnout is real, and chugging energy drinks before your 10 AM tutorial only leads to a 2 PM brain crash.",
            "What if you started your day by actually resetting your central nervous system?",
            "",
            "Sunrise Pranayama isn’t about sitting motionless like a statue. It’s about clearing brain fog, boosting blood flow, and walking into your day with laser focus while everyone else is still dragging their feet.",
            "",
            "Better energy. Sharper memory. Zero caffeine crash.",
            "",
            "⚡ CTA: Tag that one friend who needs three alarm clocks to wake up. Next session is this Saturday at 6:45 AM—meet us on the sports lawn.",
            "",
            "Hashtags: #SYCABES #StudentYogicClub #CollegeLifeHacks #NoBurnout #CampusFitness #MindfulEngineers #MindsetMatters #SunriseRoutine"
        ],
        tag="CAPTION 1B",
        bg_color="F5F2EB"
    )

    # Post 2
    add_section_heading("Post 2: Exam Stress Buster — Sound Bath & Box Breathing Workshop", level=2)
    doc.add_paragraph(
        "Visual Context: Aesthetic high-contrast graphic / reel addressing mid-sem panic, late-night cramming, and sensory overload."
    ).paragraph_format.space_after = Pt(6)

    add_callout(
        "OPTION 2A — Empathy-First & Psychological Relief Angle",
        [
            "\"That tightness in your chest right before exams isn't just 'stress.' It’s your nervous system begging for a pause.\"",
            "",
            "Mid-sem week is here. Your desk is covered in half-understood lecture notes, your screen time is spiking, and your mind is running 12 browser tabs at once.",
            "",
            "Here’s a gentle reminder: Your GPA is important, but your mental equilibrium is irreplaceable.",
            "",
            "Tomorrow evening, SYC is hosting a 40-minute Sound Bath & Guided Box Breathing Workshop. No books, no technical talk. Just guided sound resonance, deep parasympathetic down-regulation, and tools to calm exam jitters in 60 seconds.",
            "",
            "Walk in overwhelmed. Walk out anchored.",
            "",
            "🕯️ CTA: Tap the link in our bio to reserve your meditation cushion (limited to 50 students for sound acoustics). Share this with your study group—they need this more than another coffee run.",
            "",
            "Hashtags: #SYCRecruitment #ExamStressRelief #SoundBath #BoxBreathing #MentalHealthMatters #ABESCampus #MindfulStudents #BeatExamAnxiety"
        ],
        tag="CAPTION 2A",
        bg_color="EAF2EC"
    )

    add_callout(
        "OPTION 2B — Actionable 60-Second Micro-Habit Protocol",
        [
            "\"Try this 60-second breathing hack at your study desk right now (don't scroll past, just do it with us):\"",
            "",
            "1. Inhale through your nose for 4 counts. 🫁",
            "2. Hold that breath gently for 4 counts. ⏸️",
            "3. Exhale slowly through your mouth for 4 counts. 🌬️",
            "4. Hold empty for 4 counts. 🧘",
            "",
            "Repeat 3 times. Feel that sudden drop in heart rate?",
            "That’s the Box Breathing Protocol—used by top athletes, high-stress surgeons, and mindful engineers to eliminate cognitive panic in seconds.",
            "",
            "You don't have to study until your eyes burn. You study better when your mind is calm.",
            "",
            "📌 CTA: Save this post to your 'Exams' collection so you can pull it out 5 minutes before entering the exam hall. Drop a '✨' if your shoulders just dropped two inches.",
            "",
            "Hashtags: #StudyWithSYC #MicroHabits #BoxBreathing #StudentYogicClub #ExamSeasonHacks #MindfulnessInAction #CampusWellness #EngineersWhoBreathe"
        ],
        tag="CAPTION 2B",
        bg_color="F5F2EB"
    )

    # Post 3
    add_section_heading("Post 3: Behind The Scenes (BTS) — The Builders of SYC", level=2)
    doc.add_paragraph(
        "Visual Context: Candid carousel of developers writing Next.js code, designers adjusting typography, and videographers color grading reels."
    ).paragraph_format.space_after = Pt(6)

    add_callout(
        "OPTION 3A — The Contrasting Identity & Craftsmanship Angle",
        [
            "\"When people hear 'Yogic Club,' they imagine silence and stillness. They usually don't expect 5,000 lines of Next.js code and 4K cinema cameras.\"",
            "",
            "Meet the team building the modern SYC experience.",
            "Behind every sunrise session and campus wellness festival is a room full of young creators, developers, cinematographers, and organizers working in sync.",
            "",
            "We write code that handles recruitment influx. We design typography that challenges traditional college club posters. We craft films that tell authentic stories. And we do all of it while practicing the very mindfulness we advocate.",
            "",
            "We don't just run an extracurricular club—we are building an ecosystem of disciplined, conscious creators.",
            "",
            "🤝 CTA: Want to build with people who balance high ambition with deep inner calm? Our 2026–27 Recruitment cycle is open now. Link in bio to apply for Tech, Design, Video, and Operations.",
            "",
            "Hashtags: #BehindTheScenes #TeamSYC #TechAndMindfulness #StudentYogicClub #CampusLeadership #CreativeCollective #ABESEngineeringCollege"
        ],
        tag="CAPTION 3A",
        bg_color="EAF2EC"
    )

    add_callout(
        "OPTION 3B — The Relatable, Candid & Campus Life Angle",
        [
            "\"90% creative vision. 10% trying to figure out why the mic stopped recording. 100% good energy.\"",
            "",
            "A little peek behind the curtain of our media and tech team this week:",
            "- 3 drafts of the Wellness Week poster before the lead said 'make the typography breathing room larger' 🎨",
            "- Testing our recruitment portal at midnight over cold canteen samosas 💻",
            "- Laughing through bloopers while shooting our promotional reels 🎬",
            "",
            "Being in SYC isn't about being perfectly zen all the time. It’s about having a crew where you can build cool things, fail forward, laugh loudly, and grow alongside genuine people.",
            "",
            "❤️ CTA: Which slide captures your campus vibe best? 1, 2, or 3? Drop your answer below and don't forget to check out our open recruitment wings in the bio!",
            "",
            "Hashtags: #SYCCulture #BehindTheLens #CollegeClubs #CampusCreators #StudentLifeABES #SYC2026 #CreativeMinds"
        ],
        tag="CAPTION 3B",
        bg_color="F5F2EB"
    )

    # SECTION 3: PROMOTIONAL CONTENT
    add_section_heading("PART 2: PROMOTIONAL CONTENT — SYC RECRUITMENT 2026–27", level=1)
    doc.add_paragraph(
        "Target Channels: Instagram Feed / LinkedIn Long-Form / WhatsApp Society Broadcasts\n"
        "Key Objective: Demolish passive stereotypes, highlight tangible portfolio-building opportunities, and drive immediate task submissions."
    ).paragraph_format.space_after = Pt(8)

    add_callout(
        "FLAGSHIP PROMOTIONAL COPY // COHORT 2026–27",
        [
            "Most college clubs will teach you how to make noise. SYC teaches you how to build in silence—and make a lasting impact.",
            "",
            "Let’s be honest: your college years shouldn’t just be a checklist of 75% attendance, last-minute assignment submissions, and doomscrolling between lectures.",
            "",
            "You came here with ambition. But ambition without mental clarity is just burnout waiting to happen.",
            "",
            "Welcome to the Student Yogic Club (SYC) — Recruitment Cohort 2026–27.",
            "",
            "SYC is not your typical college society. We are an interdisciplinary ecosystem where inner discipline meets cutting-edge execution. We don't just sit on yoga mats; we build production-grade web systems, curate editorial magazines, direct cinematic brand films, and organize flagship wellness festivals for 1,000+ students across campus.",
            "",
            "🌿 WHAT YOU WILL ACTUALLY GAIN INSIDE SYC:",
            "",
            "1. High-Impact Portfolio Experience:",
            "   • Technical Wing: Build production web experiences with Next.js, TypeScript, and modern backend architectures.",
            "   • Graphic Design Wing: Design luxury editorial typography, campaign brand identities, and high-conversion social assets.",
            "   • Video Production Wing: Direct 4K cinematic reels, master color grading, and craft sound-designed visual narratives.",
            "   • Operations & Management: Choreograph large-scale campus events, manage sponsor relations, and direct real teams.",
            "",
            "2. The Discipline of High Performance:",
            "   • Master breathwork, mental poise, and focus rituals that keep you unshakable during 48-hour hackathons and high-stakes interviews.",
            "",
            "3. A Grounded, Ambitious Circle:",
            "   • Collaborate with the sharpest minds on campus who prioritize mental health, creative freedom, and genuine brotherhood over shallow networking.",
            "",
            "You don't need to be an expert in yoga. You just need curiosity, discipline, and a hunger to create meaningful work.",
            "",
            "⏳ DEADLINE: Wednesday, 11:59 PM",
            "📍 Open Wings: Technical • Graphic Design • Video Production • Operations & Editorial",
            "🔗 APPLY NOW: Tap the link in our bio or visit syc.abes.ac.in/apply to submit your application and domain task.",
            "",
            "Don’t just pass through college. Build yourself from the inside out."
        ],
        tag="OFFICIAL PROMOTIONAL BROADCAST",
        bg_color="EAF2EC"
    )

    # SECTION 4: REEL SCRIPT & STORYBOARD
    add_section_heading("PART 3: CREATIVE REEL SCRIPT (20–30 SECONDS)", level=1)
    p_reel_desc = doc.add_paragraph(
        "Theme: \"Why should a college student join SYC?\"\n"
        "Total Run Time: 26 Seconds  |  Pacing: Fast, energetic pattern-interrupt opening transitioning into serene, cinematic rhythm.\n"
        "Audio Architecture: Begins with chaotic muffled alarm & typing noises, abruptly cutting into a crisp Lo-Fi / Fred Again-style ambient electronic beat."
    )
    p_reel_desc.paragraph_format.space_after = Pt(10)

    # Table for Script
    tbl_script = doc.add_table(rows=6, cols=5)
    tbl_script.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_script.autofit = False

    col_widths = [Inches(1.0), Inches(1.5), Inches(1.4), Inches(1.7), Inches(1.2)]
    headers = ["TIMECODE", "VISUAL / CAMERA", "ON-SCREEN TEXT", "VOICEOVER (V.O.)", "SFX & MUSIC"]

    for col_idx, text in enumerate(headers):
        cell = tbl_script.cell(0, col_idx)
        cell.width = col_widths[col_idx]
        set_cell_background(cell, "1C1D1A")
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        p = cell.paragraphs[0]
        r = p.add_run(text)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    script_rows = [
        (
            "0:00 – 0:03\n[The Hook]",
            "Fast, dizzying montage: Glowing phone screen at 2 AM, laptop with 20 tabs, anxious finger tapping in exam hall. Close-up of eyes opening.",
            "\"Be honest...\"\nWHEN WAS THE LAST TIME YOUR MIND WAS ACTUALLY QUIET?",
            "(V.O. - Relatable, direct):\n\"Be honest. When was the last time your mind was actually quiet?\"",
            "Muffled alarm beeps + typing noise abruptly cuts with a sharp [WHOOSH] into dead silence."
        ),
        (
            "0:03 – 0:08\n[Stereotype\nBust]",
            "Split-second cliché stock meditation image shattered by digital glitch. Cut to rapid cuts: dev pushing git commit, camera gimbal, Figma typography.",
            "(And no, this isn't just about sitting in a park.)",
            "\"And no... we’re not asking you to become a monk or wake up at 4 AM to stare at a wall.\"",
            "Digital glitch [STATIC SNAP], followed by a deep ambient bass drop and crisp drums."
        ),
        (
            "0:08 – 0:17\n[The Value &\nReality]",
            "Golden hour cinema shot: 40 students breathing together on lawn in sync, then cutting to high-energy team brainstorming, event cheering.",
            "Where Code meets Calm.\nWhere Design meets Discipline.",
            "\"Welcome to SYC. Where engineering students learn to build real projects, shoot cinematic films, and master the mental calm that prevents college burnout.\"",
            "Warm melodic bassline grooves smoothly; layered sound of collective deep breath."
        ),
        (
            "0:17 – 0:22\n[The Trans-\nformation]",
            "Medium shot of student smiling, closing laptop with confidence, looking energized and laughing with teammates in sun.",
            "Real Portfolio.\nUnshakable Focus.\nYour Sanctuary.",
            "\"You don't just gain line items for your resume. You build a brain that doesn’t panic under pressure.\"",
            "Uplifting harmonic chord builds to an inspiring crescendo."
        ),
        (
            "0:22 – 0:26\n[Ending / CTA]",
            "Clean typography animation over dark luxury minimalist background: Stylized SYC letters, recruitment deadlines, and wings.",
            "SYC RECRUITMENT 2026–27\nCloses Wednesday 11:59 PM\nLINK IN BIO",
            "\"Don’t just survive your degree. Master yourself. SYC Recruitment is live. Link in bio.\"",
            "Final crisp beat tap [PING], ambient reverb trails off cleanly."
        )
    ]

    for row_idx, data in enumerate(script_rows, start=1):
        bg = "FBF9F5" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = tbl_script.cell(row_idx, col_idx)
            cell.width = col_widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x2A, 0x2B, 0x28)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x2D, 0x4A, 0x3E)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # SECTION 5: STRATEGIC RATIONALE & INTERVIEW DEFENSE
    add_section_heading("PART 4: STRATEGIC RATIONALE & INTERVIEW DEFENSE", level=1)
    doc.add_paragraph(
        "Below are structured, articulate answers to the specific interview questions outlined in the recruitment prompt:"
    ).paragraph_format.space_after = Pt(6)

    add_callout(
        "QUESTION 1: WHY DID YOU CHOOSE YOUR WRITING STYLE?",
        [
            "\"I chose an 'Empathetic & Aspirational' tone that bridges psychological honesty with modern design sensibility.",
            "",
            "1. Rejection of Spiritual Clichés: Traditional wellness copy is often perceived as passive or detached from real student struggles. By talking about 8-hour lectures, DSA prep, and late-night screen time, the copy immediately establishes peer-level trust.",
            "2. Cadence & Rhythm Modulation: Great copywriting reads like spoken poetry. I structure sentences with variable lengths: short 5-word declarations to arrest attention, followed by rhythmic paragraphs for substance, and zero-friction calls to action.",
            "3. Professional Polish: It positions SYC as an elite, high-agency society—appealing equally to serious tech builders, artistic designers, and thoughtful students.\""
        ],
        tag="INTERVIEW DEFENSE #1",
        bg_color="EAF2EC"
    )

    add_callout(
        "QUESTION 2: HOW DID YOU DEVELOP THE HOOKS?",
        [
            "\"I utilized three proven psychological frameworks across the deliverables:",
            "",
            "• The Pattern Interrupt (Reel Script): Most college reels start with loud music or generic slogans. Starting with 'When was the last time your mind was actually quiet?' mirrors the exact mental state of a student doomscrolling in bed at 1 AM, instantly freezing their thumb.",
            "• The Paradox Frame (BTS Post 3A): Pairing two contradictory ideas ('Yogic Club' vs. '5,000 lines of Next.js code') creates cognitive dissonance. The human brain is wired to read further to resolve the conflict.",
            "• The Micro-Optimization Frame (Post 1B): Reframing Pranayama as a 'cheat code' rather than a ritual leverages the engineering mindset of hacking daily performance.\""
        ],
        tag="INTERVIEW DEFENSE #2",
        bg_color="F5F2EB"
    )

    add_callout(
        "QUESTION 3: HOW WOULD YOU ADAPT CONTENT FOR DIFFERENT AUDIENCES?",
        [
            "\"Audience segmentation is critical for campus engagement:",
            "",
            "• Freshers (1st Years): They feel overwhelmed by campus transition and peer pressure. Content should focus on belonging, warmth, mental grounding, and discovering new creative wings without judgment.",
            "• Sophomores & Pre-Final Years (2nd & 3rd Years): They are hyper-focused on careers, hackathons, and resumes. Content must highlight tangible portfolio deliverables (building real web apps, directing brand films, leading teams) alongside mental resilience under stress.",
            "• Tech Enthusiasts vs. Creative Designers: For tech candidates, frame SYC through execution discipline and systems architecture. For design/media candidates, highlight editorial aesthetics, typography freedom, and cinematic camera equipment.\""
        ],
        tag="INTERVIEW DEFENSE #3",
        bg_color="EAF2EC"
    )

    add_callout(
        "QUESTION 4: HOW DO YOU MAKE A CAPTION MORE ENGAGING?",
        [
            "\"I adhere to the 4-part 'H.V.F.T.' architecture:",
            "",
            "1. Hook (Top 12 Words): Must resolve before Instagram's '...more' fold. If line 1 doesn't create an open loop, the rest doesn't exist.",
            "2. Value (The Paradigm Shift): Offer an actionable insight, emotional relief, or an unconventional perspective.",
            "3. Formatting (Visual Breathing Room): Break paragraphs into 1-2 sentence chunks with selective emojis as visual bullet signposts.",
            "4. Trigger (Low-Friction CTA): Avoid high-effort asks like 'Click the link and fill a 5-page form.' Start with micro-engagements: 'Save this for tomorrow morning' or 'Tag your late-night study buddy.'\""
        ],
        tag="INTERVIEW DEFENSE #4",
        bg_color="F5F2EB"
    )

    # SECTION 6: PREVIOUS WORK & PORTFOLIO
    add_section_heading("PART 5: PREVIOUS WORK & PORTFOLIO SAMPLES", level=1)
    
    add_callout(
        "PORTFOLIO EXCERPT 1 — EDITORIAL ESSAY ON TECH & MINDFULNESS",
        [
            "Title: The Mindful Engineer: Why Breathwork Is Biological Garbage Collection",
            "\"We teach computer science students how to optimize time complexity O(n), but we never teach them how to regulate cognitive overload when their code breaks in production at 3 AM. A compiler doesn't care about your panic; it only reflects your clarity. Breathwork isn't mysticism—it is biological garbage collection for your nervous system.\""
        ],
        tag="WRITING SAMPLE 1",
        bg_color="EAF2EC"
    )

    add_callout(
        "PORTFOLIO EXCERPT 2 — CAMPUS VIRAL EVENT COPY",
        [
            "Title: The 72-Hour Digital Detox Challenge",
            "\"Your phone told you your screen time was 7 hours and 42 minutes yesterday. That’s an entire working day spent consuming other people’s lives. Starting this Friday, SYC is hosting the 72-Hour Campus Detox: No doomscrolling after 9 PM, 15 minutes of lawn walking, and three offline group circles. Are you brave enough to disconnect to reconnect?\""
        ],
        tag="WRITING SAMPLE 2",
        bg_color="F5F2EB"
    )

    add_callout(
        "PORTFOLIO EXCERPT 3 — PRODUCT LAUNCH MICRO-COPY",
        [
            "Context: Official Launch of SYC's 2026 Recruitment Platform",
            "\"Fast. Minimal. Intentional. We rebuilt the recruitment portal from the ground up: zero clutter, sub-100ms transitions, and an applicant tracking experience designed to respect your time. Go see what mindful software feels like at syc.abes.ac.in.\""
        ],
        tag="WRITING SAMPLE 3",
        bg_color="EAF2EC"
    )

    # Sign-off
    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(16)
    r_end = p_end.add_run("Submitted with dedication, creative rigor, and respect for the SYC vision.\nTeam SYC // Recruitment Cohort 2026–27")
    r_end.font.bold = True
    r_end.font.size = Pt(10)
    r_end.font.color.rgb = RGBColor(0x2D, 0x4A, 0x3E)

    output_path = os.path.join(os.getcwd(), "CONTENT_TASK_SUBMISSION.docx")
    doc.save(output_path)
    print(f"Successfully generated Word document at: {output_path}")

if __name__ == "__main__":
    create_document()
